"""Attachment ingestion for chat and council mode.

Extracts text from PDF, DOCX, text/code files, spreadsheets, and folders.
Images are returned separately so they can be routed to a vision-capable model.
Large text is chunked and summarized to keep context comprehensive but bounded.
"""

import logging
from pathlib import Path

logger = logging.getLogger(__name__)

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".py", ".js", ".ts", ".java", ".c", ".cpp",
    ".h", ".cs", ".go", ".rs", ".rb", ".php", ".sh", ".bat", ".ps1", ".sql",
    ".json", ".xml", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".log",
    ".html", ".htm", ".css",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
SHEET_EXTS = {".xlsx", ".xls", ".csv"}


def is_image(path) -> bool:
    return Path(path).suffix.lower() in IMAGE_EXTS


def collect_paths(paths):
    """Expand folders recursively into a list of file paths."""
    out = []
    for p in paths:
        path = Path(p)
        if path.is_dir():
            for child in sorted(path.rglob("*")):
                if child.is_file():
                    out.append(child)
        elif path.is_file():
            out.append(path)
    return out


def extract_text(path, max_chars: int = 20000) -> str:
    """Extract text from a single file (best effort)."""
    path = Path(path)
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            return _extract_pdf(path, max_chars)
        if ext == ".docx":
            return _extract_docx(path, max_chars)
        if ext in SHEET_EXTS:
            return _extract_sheet(path, max_chars)
        if ext in IMAGE_EXTS:
            return ""
        if ext in TEXT_EXTS or ext == "":
            return path.read_text(encoding="utf-8", errors="ignore")[:max_chars]
        # Unknown: try as text
        return path.read_text(encoding="utf-8", errors="ignore")[:max_chars]
    except Exception as e:
        logger.exception("Failed to read %s", path)
        return f"(Error reading {path.name}: {e})"


def _extract_pdf(path, max_chars):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    parts = []
    total = 0
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        parts.append(text)
        total += len(text)
        if total >= max_chars:
            break
    return "\n".join(parts)[:max_chars]


def _extract_docx(path, max_chars):
    import docx
    doc = docx.Document(str(path))
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)[:max_chars]


def _extract_sheet(path, max_chars):
    import pandas as pd
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path)
        return df.head(200).to_string(index=False)[:max_chars]
    sheets = pd.read_excel(path, sheet_name=None)
    parts = []
    for name, df in sheets.items():
        if df is None or df.empty:
            continue
        parts.append(f"## Sheet: {name}")
        parts.append(df.head(100).to_string(index=False))
    return "\n\n".join(parts)[:max_chars]


def chunk_text(text: str, max_chars: int = 8000, overlap: int = 200):
    """Split text into overlapping chunks."""
    if len(text) <= max_chars:
        return [text]
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + max_chars, len(text))
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = end - overlap
    return chunks


def summarize_large(text: str, summarize_fn=None, max_total: int = 12000) -> str:
    """Chunk + summarize long text so the final context stays bounded."""
    if len(text) <= max_total:
        return text
    if summarize_fn is None:
        return text[:max_total]
    parts = []
    for chunk in chunk_text(text, max_chars=8000):
        try:
            parts.append(summarize_fn(
                "Summarize the following content into concise bullet points, "
                "preserving key facts, numbers, and structure:\n\n" + chunk
            ))
        except Exception:
            parts.append(chunk[:2000])
    combined = "\n\n".join(parts)
    return combined[:max_total] if len(combined) > max_total else combined


def build_attachment_context(paths, summarize_fn=None, max_total: int = 12000):
    """Returns (text_context, image_paths) for a set of attachments."""
    files = collect_paths(paths or [])
    text_parts = []
    images = []
    for f in files:
        if is_image(f):
            images.append(str(f))
            continue
        text = extract_text(f)
        if text.strip():
            text_parts.append(f"## Attachment: {f.name}\n{text}")
    combined = "\n\n".join(text_parts)
    combined = summarize_large(combined, summarize_fn, max_total)
    return combined, images
